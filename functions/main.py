import os
import functions_framework
from google.cloud import firestore, storage
import docx
import PyPDF2
import requests

# Environment Variables
TRANSLATION_API_URL = os.environ.get("TRANSLATION_API_URL", "https://ml-backend.internal/translate")
USE_EXTERNAL_LLM = os.environ.get("USE_EXTERNAL_LLM", "false").lower() == "true"
OPENAI_API_KEY = os.environ.get("OPENAI_API_KEY")

# Initialize clients
db = firestore.Client()
storage_client = storage.Client()

@functions_framework.cloud_event
def process_translation_task(cloud_event):
    """
    Triggered by a new document in the 'translationTasks' collection.
    Orchestrates the document translation workflow.
    """
    task_id = cloud_event.params['taskId']
    print(f"Processing task ID: {task_id}")

    # 1. Get task data from Firestore
    task_ref = db.collection('translationTasks').document(task_id)
    task_data = task_ref.get().to_dict()
    if not task_data:
        print(f"Error: Task {task_id} not found in Firestore.")
        return

    update_task_status(task_ref, 'processing', 10)

    # 2. Download original file from Cloud Storage
    bucket_name = os.environ.get('GCS_BUCKET') # You need to set this env var
    file_name = task_data['fileName']
    source_blob_path = f"uploads/{task_id}/{file_name}"
    
    try:
        source_blob = storage_client.bucket(bucket_name).blob(source_blob_path)
        if not source_blob.exists():
             raise FileNotFoundError(f"Source file not found at {source_blob_path}")
        
        temp_file_path = f"/tmp/{file_name}"
        source_blob.download_to_filename(temp_file_path)
        print(f"Downloaded {file_name} to {temp_file_path}")
        update_task_status(task_ref, 'processing', 25)
    except Exception as e:
        print(f"Error downloading file: {e}")
        update_task_status(task_ref, 'failed', error_message=str(e))
        return

    # 3. Extract text from the document
    try:
        text_content = extract_text(temp_file_path, file_name)
        update_task_status(task_ref, 'processing', 50)
    except Exception as e:
        print(f"Error extracting text: {e}")
        update_task_status(task_ref, 'failed', error_message=f"Text extraction failed: {e}")
        return

    # 4. Translate text (stubbed)
    try:
        translated_text = translate_text(
            text_content,
            task_data['srcLang'],
            task_data['tgtLang']
        )
        update_task_status(task_ref, 'processing', 75)
    except Exception as e:
        print(f"Error during translation: {e}")
        update_task_status(task_ref, 'failed', error_message=f"Translation failed: {e}")
        return
        
    # 5. Generate output files (stubbed for TXT)
    try:
        output_paths = generate_outputs(
            task_id,
            translated_text,
            bucket_name
        )
        update_task_status(task_ref, 'completed', 100, outputs=output_paths)
    except Exception as e:
        print(f"Error generating output files: {e}")
        update_task_status(task_ref, 'failed', error_message=f"Output generation failed: {e}")
        return
        
    print(f"Task {task_id} completed successfully.")

def update_task_status(task_ref, status, progress, outputs=None, error_message=None):
    """Updates the task document in Firestore."""
    update_data = {
        'status': status,
        'progress': progress,
        'updatedAt': firestore.SERVER_TIMESTAMP
    }
    if outputs:
        update_data['outputs'] = outputs
    if error_message:
        update_data['errors'] = firestore.ArrayUnion([error_message])
    task_ref.update(update_data)
    print(f"Updated task {task_ref.id} to status: {status}, progress: {progress}%")


def extract_text(file_path, file_name):
    """Extracts text from various document types."""
    _, extension = os.path.splitext(file_name.lower())
    if extension == '.pdf':
        return extract_text_from_pdf(file_path)
    elif extension == '.docx':
        return extract_text_from_docx(file_path)
    elif extension == '.txt':
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    else:
        # Fallback to OCR could be implemented here
        raise ValueError(f"Unsupported file type: {extension}")

def extract_text_from_pdf(file_path):
    """Extracts text from a PDF file."""
    text = ""
    with open(file_path, 'rb') as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            text += page.extract_text() or "" # Add empty string if None
    # If text is empty, consider OCR fallback, e.g., using pytesseract
    if not text.strip():
        print("Warning: PDF text extraction yielded empty result. OCR fallback needed.")
    return text

def extract_text_from_docx(file_path):
    """Extracts text from a DOCX file."""
    doc = docx.Document(file_path)
    return "\n".join([para.text for para in doc.paragraphs])

def translate_text(text, src_lang, tgt_lang):
    """
    Translates text using either an internal API or an external LLM.
    This is a STUB function.
    """
    print(f"Translating text from {src_lang} to {tgt_lang}...")
    if USE_EXTERNAL_LLM:
        # Stub for OpenAI API call
        print("Using external LLM (OpenAI stub)")
        # response = requests.post("https://api.openai.com/v1/chat/completions", ... )
        # return response.json()['choices'][0]['message']['content']
        return f"[Mock OpenAI Translation]: {text}"
    else:
        # Stub for internal FastAPI translation service
        print(f"Using internal translation service at {TRANSLATION_API_URL}")
        # response = requests.post(TRANSLATION_API_URL, json={
        #     "text": text, "source_lang": src_lang, "target_lang": tgt_lang
        # })
        # response.raise_for_status()
        # return response.json()['translation']
        return f"[Mock Internal Translation]: {text}"

def generate_outputs(task_id, translated_text, bucket_name):
    """Generates output files and uploads them to Storage."""
    output_paths = {}
    
    # Generate TXT (fallback)
    txt_file_path = f"/tmp/{task_id}_output.txt"
    with open(txt_file_path, "w", encoding="utf-8") as f:
        f.write(translated_text)
    
    txt_blob_path = f"results/{task_id}/output.txt"
    upload_blob(bucket_name, txt_file_path, txt_blob_path)
    output_paths["txt"] = txt_blob_path
    
    # Stub for DOCX generation
    # doc = docx.Document()
    # doc.add_paragraph(translated_text)
    # docx_file_path = f"/tmp/{task_id}_output.docx"
    # doc.save(docx_file_path)
    # docx_blob_path = f"results/{task_id}/output.docx"
    # upload_blob(bucket_name, docx_file_path, docx_blob_path)
    # output_paths["docx"] = docx_blob_path

    # Stub for PDF generation
    # ... logic to convert HTML to PDF ...
    
    return output_paths

def upload_blob(bucket_name, source_file_name, destination_blob_name):
    """Uploads a file to the bucket."""
    bucket = storage_client.bucket(bucket_name)
    blob = bucket.blob(destination_blob_name)
    blob.upload_from_filename(source_file_name)
    print(f"File {source_file_name} uploaded to {destination_blob_name}.")
